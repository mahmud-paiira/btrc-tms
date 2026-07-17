from django.utils import timezone
from datetime import date
import csv
import io

import openpyxl
from apps.common.utils import to_english_digits
from django.db.models import Q, Count, Prefetch, Subquery
from django.http import HttpResponse
from rest_framework import viewsets, status, permissions, filters as drf_filters
from rest_framework.decorators import action
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.response import Response
from django_filters import rest_framework as filters
from django_filters.rest_framework import DjangoFilterBackend

from .models import Trainer, TrainerMapping
from .serializers import (
    TrainerListSerializer, TrainerDetailSerializer,
    TrainerWriteSerializer, TrainerMappingSerializer,
    TrainerApprovalSerializer,
)
from apps.centers.models import ActionLog


class IsHeadOffice(permissions.BasePermission):
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        return request.user.user_type == 'head_office' or request.user.is_superuser


class TrainerFilter(filters.FilterSet):
    mapping_center = filters.NumberFilter(field_name='mappings__center_id', label='Center')

    class Meta:
        model = Trainer
        fields = ('status', 'approval_status', 'expertise_area', 'years_of_experience', 'mapping_center')


class HOTrainerViewSet(viewsets.ModelViewSet):
    queryset = Trainer.objects.select_related('user', 'approved_by').prefetch_related(
        Prefetch('mappings', queryset=TrainerMapping.objects.select_related('center', 'course', 'approved_by')),
    ).all()
    permission_classes = [permissions.IsAuthenticated, IsHeadOffice]
    filter_backends = (DjangoFilterBackend, drf_filters.SearchFilter, drf_filters.OrderingFilter)
    filterset_class = TrainerFilter
    search_fields = (
        'trainer_no', 'nid', 'birth_certificate_no',
        'user__email', 'user__phone',
        'user__full_name_bn', 'user__full_name_en',
    )
    ordering_fields = ('trainer_no', 'years_of_experience', 'created_at')
    ordering = ('-created_at',)

    def get_serializer_class(self):
        if self.action == 'list':
            return TrainerListSerializer
        if self.action in ('create', 'update', 'partial_update'):
            return TrainerWriteSerializer
        return TrainerDetailSerializer

    def _log(self, request, action_desc, target_id=None):
        ActionLog.objects.create(
            user=request.user,
            action=action_desc,
            target_type='Trainer',
            target_id=str(target_id) if target_id else '',
            description=action_desc,
            ip_address=request.META.get('REMOTE_ADDR', ''),
        )

    def perform_create(self, serializer):
        trainer = serializer.save()
        self._log(self.request, f'Created trainer {trainer.trainer_no}', trainer.id)

    def perform_update(self, serializer):
        trainer = serializer.save()
        self._log(self.request, f'Updated trainer {trainer.trainer_no}', trainer.id)

    def perform_destroy(self, instance):
        self._log(self.request, f'Deleted trainer {instance.trainer_no}', instance.id)
        instance.delete()

    @action(detail=False, methods=['post'], url_path='bulk_delete')
    def bulk_delete(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return Response({'error': 'কোন আইডি প্রদান করা হয়নি'}, status=400)
        deleted = 0
        errors = []
        for pk in ids:
            try:
                obj = self.get_queryset().get(pk=pk)
                self.perform_destroy(obj)
                deleted += 1
            except Exception as e:
                msg = str(e.detail[0]) if hasattr(e, 'detail') and isinstance(e.detail, list) else str(e)
                errors.append(msg)
        return Response({'deleted': deleted, 'errors': errors})

    @action(detail=False, methods=['post'], parser_classes=[MultiPartParser, FormParser], url_path='import_list')
    def import_list(self, request):
        from apps.centers.models import Center
        from apps.courses.models import Course
        from apps.accounts.models import User

        file = request.FILES.get('file')
        if not file:
            return Response({'error': 'ফাইল নির্বাচন করুন'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            rows_iter = iter(ws.iter_rows(values_only=True))
            header_row = [str(c).strip().lower() if c is not None else '' for c in next(rows_iter)]
        except Exception:
            file.seek(0)
            try:
                content = file.read().decode('utf-8-sig')
                reader = csv.DictReader(io.StringIO(content))
                header_row = [h.strip().lower() for h in reader.fieldnames]
                rows_iter = reader
            except Exception:
                return Response({'error': 'ভুল ফাইল ফরম্যাট। Excel (.xlsx) বা CSV ফাইল আপলোড করুন।'}, status=400)

        bn_required = {'প্রশিক্ষক নং'}
        header_set = set(header_row)
        if not bn_required.issubset(header_set):
            return Response({
                'error': f'প্রয়োজনীয় কলাম নেই। হেডারে "প্রশিক্ষক নং" থাকা আবশ্যক।',
                'detected_headers': header_row,
            }, status=400)

        field_map = {
            'প্রশিক্ষক নং': 'trainer_no', 'trainer_no': 'trainer_no',
            'নাম (বাংলা)': 'full_name_bn', 'name_bn': 'full_name_bn',
            'নাম (ইংরেজি)': 'full_name_en', 'name_en': 'full_name_en',
            'ইমেইল': 'email', 'email': 'email',
            'ফোন': 'phone', 'phone': 'phone',
            'এনআইডি': 'nid', 'nid': 'nid',
            'শিক্ষাগত যোগ্যতা': 'education_qualification',
            'অভিজ্ঞতা (বছর)': 'years_of_experience',
            'দক্ষতা': 'expertise_area', 'দক্ষতার ক্ষেত্র': 'expertise_area', 'expertise_area': 'expertise_area',
            'স্ট্যাটাস': 'status', 'status': 'status',
            'অনুমোদন স্ট্যাটাস': 'approval_status', 'approval_status': 'approval_status',
            'কেন্দ্রের কোড': 'center_code', 'center_code': 'center_code',
            'কেন্দ্রের নাম': 'center_name', 'center_name': 'center_name',
            'কোর্সের কোড': 'course_code', 'course_code': 'course_code',
            'কোর্সের নাম': 'course_name', 'course_name': 'course_name',
        }

        results = {'created': 0, 'updated': 0, 'errors': []}
        for row_idx, row in enumerate(rows_iter, start=2):
            try:
                if isinstance(row, dict):
                    raw = row
                else:
                    if all(c is None for c in row):
                        continue
                    raw = dict(zip(header_row, [str(c).strip() if c is not None else '' for c in row]))

                data = {}
                for k, v in raw.items():
                    mapped = field_map.get(k.strip().lower(), k.strip().lower())
                    data[mapped] = v.strip() if v else ''

                if data.get('phone'):
                    data['phone'] = to_english_digits(data['phone']).replace('-', '')
                if data.get('nid'):
                    data['nid'] = to_english_digits(data['nid']).replace(' ', '').replace('-', '')
                if data.get('email') and data['email'] in ('\u2013', '\u2014', '-', '/'):
                    data['email'] = ''

                trainer_no = data.get('trainer_no', '').strip()
                if not trainer_no:
                    results['errors'].append(f'সারি {row_idx}: প্রশিক্ষক নং আবশ্যক')
                    continue

                existing = Trainer.objects.filter(trainer_no=trainer_no).first()
                if trainer_no.isdigit() and not existing:
                    generated_no = f'BRW-{trainer_no}'
                    existing = Trainer.objects.filter(trainer_no=generated_no).first()
                    nid = data.get('nid', '')
                    if not existing:
                        if not nid:
                            results['errors'].append(f'সারি {row_idx}: "{trainer_no}" - এনআইডি ছাড়া নতুন প্রশিক্ষক তৈরি সম্ভব নয়')
                            continue
                        existing_user = User.objects.filter(nid=nid).first() if nid else None
                        if existing_user:
                            existing = Trainer.objects.filter(user=existing_user).first()
                            if not existing:
                                existing = Trainer.objects.create(user=existing_user, trainer_no=generated_no, nid=nid, date_of_birth=date.today())
                            existing.trainer_no = generated_no
                        if not existing or not existing.pk:
                            phone = data.get('phone', '')
                            if phone and User.objects.filter(phone=phone).exists():
                                phone = ''
                            email = data.get('email', '') or f'{generated_no.lower()}@brtc.app'
                            user = User.objects.create_user(
                                email=email,
                                password='trainer@123',
                                full_name_bn=data.get('full_name_bn', '') or '\u2014',
                                full_name_en=data.get('full_name_en', '') or '\u2014',
                                phone=phone,
                                nid=nid,
                                user_type='trainer',
                            )
                            existing = Trainer.objects.create(
                                user=user,
                                trainer_no=generated_no,
                                nid=nid,
                                date_of_birth=date.today(),
                                education_qualification=data.get('education_qualification', ''),
                                years_of_experience=int(data['years_of_experience']) if data.get('years_of_experience') else 0,
                                expertise_area=data.get('expertise_area', ''),
                                status=data.get('status', 'pending'),
                                approval_status=data.get('approval_status', 'pending'),
                            )
                    trainer_no = generated_no
                if not existing:
                    results['errors'].append(f'সারি {row_idx}: "{trainer_no}" প্রশিক্ষক নং পাওয়া যায়নি')
                    continue
                profile = existing.user
                if data.get('full_name_bn'):
                    profile.full_name_bn = data['full_name_bn']
                if data.get('full_name_en'):
                    profile.full_name_en = data['full_name_en']
                if data.get('phone'):
                    profile.phone = data['phone']
                if data.get('email'):
                    profile.email = data['email']
                profile.save()
                if data.get('education_qualification') is not None:
                    existing.education_qualification = data.get('education_qualification') or ''
                if data.get('years_of_experience') is not None:
                    existing.years_of_experience = int(data['years_of_experience']) if data['years_of_experience'] else None
                if data.get('expertise_area') is not None:
                    existing.expertise_area = data.get('expertise_area') or ''
                if data.get('status') is not None:
                    existing.status = data['status']
                if data.get('approval_status') is not None:
                    existing.approval_status = data['approval_status']
                existing.save()
                results['updated'] += 1

                center = None
                center_code = data.get('center_code', '').strip()
                center_name = data.get('center_name', '').strip()
                if center_code:
                    center = Center.objects.filter(code__iexact=center_code).first()
                if not center and center_name:
                    center = Center.objects.filter(name_bn=center_name).first()
                if center:
                    course = None
                    course_code = data.get('course_code', '').strip()
                    course_name = data.get('course_name', '').strip()
                    if course_code:
                        course = Course.objects.filter(code__iexact=course_code).first()
                    if not course and course_name:
                        course = Course.objects.filter(name_bn=course_name).first()
                    TrainerMapping.objects.update_or_create(
                        trainer=existing,
                        center=center,
                        defaults={'course': course, 'status': 'active'},
                    )
                    if not profile.center:
                        profile.center = center
                        profile.save(update_fields=['center'])
            except Exception as e:
                results['errors'].append(f'সারি {row_idx}: {str(e)}')

        return Response(results)

    @action(detail=False, methods=['get'], url_path='download_template')
    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Template'
        headers = ['প্রশিক্ষক নং', 'নাম (বাংলা)', 'নাম (ইংরেজি)', 'ইমেইল', 'ফোন',
                   'এনআইডি', 'শিক্ষাগত যোগ্যতা', 'অভিজ্ঞতা (বছর)', 'দক্ষতার ক্ষেত্র',
                   'স্ট্যাটাস', 'অনুমোদন স্ট্যাটাস',
                   'কেন্দ্রের কোড', 'কেন্দ্রের নাম', 'কোর্সের কোড']
        ws.append(headers)
        sample = ['', 'উদাহরণ নাম', 'Example Name', 'email@example.com', '০১৭XXXXXXXX',
                  '', '', '', '', 'pending', 'pending', 'RSH_TCU', 'রাজশাহী ট্রেনিং সেন্টার', 'DTP-0001']
        ws.append(sample)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = openpyxl.styles.Font(bold=True)
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="trainer_import_template.xlsx"'
        wb.save(response)
        return response

    @action(detail=False)
    def pending(self, request):
        qs = self.queryset.filter(approval_status=Trainer.ApprovalStatus.PENDING)
        page = self.paginate_queryset(qs)
        if page is not None:
            serializer = TrainerDetailSerializer(page, many=True)
            return self.get_paginated_response(serializer.data)
        serializer = TrainerDetailSerializer(qs, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        trainer = self.get_object()
        serializer = TrainerApprovalSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        trainer.approval_status = Trainer.ApprovalStatus.APPROVED
        trainer.status = Trainer.Status.ACTIVE
        trainer.approved_by = request.user
        trainer.approved_at = timezone.now()
        trainer.save(update_fields=['approval_status', 'status', 'approved_by', 'approved_at'])
        self._log(request, f'Approved trainer {trainer.trainer_no}', trainer.id)
        return Response(TrainerDetailSerializer(trainer).data)

    @action(detail=True, methods=['post'])
    def reject(self, request, pk=None):
        trainer = self.get_object()
        remarks = request.data.get('remarks', '')
        trainer.approval_status = Trainer.ApprovalStatus.REJECTED
        trainer.status = Trainer.Status.PENDING
        trainer.approved_by = request.user
        trainer.approved_at = timezone.now()
        trainer.save(update_fields=['approval_status', 'status', 'approved_by', 'approved_at'])
        self._log(request, f'Rejected trainer {trainer.trainer_no}: {remarks}', trainer.id)
        return Response(TrainerDetailSerializer(trainer).data)

    @action(detail=True, methods=['post'])
    def suspend(self, request, pk=None):
        trainer = self.get_object()
        trainer.status = Trainer.Status.SUSPENDED
        trainer.save(update_fields=['status'])
        self._log(request, f'Suspended trainer {trainer.trainer_no}', trainer.id)
        return Response(TrainerDetailSerializer(trainer).data)

    @action(detail=True, methods=['post'])
    def activate(self, request, pk=None):
        trainer = self.get_object()
        if trainer.approval_status != Trainer.ApprovalStatus.APPROVED:
            return Response({'error': 'প্রথমে প্রশিক্ষককে অনুমোদন করুন'}, status=400)
        trainer.status = Trainer.Status.ACTIVE
        trainer.save(update_fields=['status'])
        self._log(request, f'Activated trainer {trainer.trainer_no}', trainer.id)
        return Response(TrainerDetailSerializer(trainer).data)

    @action(detail=True, methods=['post'])
    def blacklist(self, request, pk=None):
        trainer = self.get_object()
        trainer.status = Trainer.Status.SUSPENDED
        trainer.approval_status = Trainer.ApprovalStatus.REJECTED
        trainer.save(update_fields=['status', 'approval_status'])
        self._log(request, f'Blacklisted trainer {trainer.trainer_no}', trainer.id)
        return Response({'status': 'blacklisted', 'trainer_no': trainer.trainer_no})

    @action(detail=True)
    def track(self, request, pk=None):
        trainer = self.get_object()
        return Response(TrainerDetailSerializer(trainer).data)

    @action(detail=False, methods=['get'])
    def track_by_number(self, request):
        trainer_no = request.query_params.get('trainer_no', '')
        try:
            trainer = self.queryset.get(trainer_no__iexact=trainer_no)
            return Response(TrainerDetailSerializer(trainer).data)
        except Trainer.DoesNotExist:
            return Response({'error': 'প্রশিক্ষক পাওয়া যায়নি'}, status=404)

    @action(detail=False, methods=['get'])
    def track_by_nid(self, request):
        nid = request.query_params.get('nid', '')
        try:
            trainer = self.queryset.get(nid__iexact=nid)
            return Response(TrainerDetailSerializer(trainer).data)
        except Trainer.DoesNotExist:
            return Response({'error': 'প্রশিক্ষক পাওয়া যায়নি'}, status=404)

    @action(detail=False, methods=['get'])
    def track_by_mobile(self, request):
        phone = request.query_params.get('phone', '')
        try:
            trainer = self.queryset.get(user__phone__iexact=phone)
            return Response(TrainerDetailSerializer(trainer).data)
        except Trainer.DoesNotExist:
            return Response({'error': 'প্রশিক্ষক পাওয়া যায়নি'}, status=404)

    @action(detail=False, methods=['get'])
    def track_by_bcn(self, request):
        bcn = request.query_params.get('bcn', '')
        try:
            trainer = self.queryset.get(birth_certificate_no__iexact=bcn)
            return Response(TrainerDetailSerializer(trainer).data)
        except Trainer.DoesNotExist:
            return Response({'error': 'প্রশিক্ষক পাওয়া যায়নি'}, status=404)

    @action(detail=False, methods=['post'])
    def map(self, request):
        mapping_serializer = TrainerMappingSerializer(data=request.data)
        mapping_serializer.is_valid(raise_exception=True)
        mapping = mapping_serializer.save()
        self._log(request, f'Mapped trainer {mapping.trainer.trainer_no} to {mapping.center.code} / {mapping.course.code}')
        return Response(TrainerMappingSerializer(mapping).data, status=201)

    @action(detail=True, methods=['post'])
    def unmap(self, request, pk=None):
        mapping_id = request.data.get('mapping_id')
        if not mapping_id:
            return Response({'error': 'mapping_id required'}, status=400)
        try:
            mapping = TrainerMapping.objects.get(id=mapping_id, trainer_id=pk)
            mapping.delete()
            self._log(request, f'Unmapped trainer mapping {mapping_id}')
            return Response({'status': 'unmapped'})
        except TrainerMapping.DoesNotExist:
            return Response({'error': 'Mapping not found'}, status=404)

    @action(detail=False, methods=['get'])
    def export(self, request):
        qs = self.filter_queryset(self.queryset)
        data = []
        for t in qs:
            center_names = ', '.join(
                dict.fromkeys(m.center.name_bn for m in t.mappings.all() if m.center)
            ) if t.mappings.all() else (
                t.user.center.name_bn if t.user and getattr(t.user, 'center', None) else ''
            )
            data.append({
                'trainer_no': t.trainer_no,
                'name_bn': t.user.full_name_bn,
                'name_en': t.user.full_name_en,
                'email': t.user.email,
                'phone': t.user.phone,
                'nid': t.nid,
                'center': center_names,
                'status': t.status,
                'approval_status': t.approval_status,
                'expertise_area': t.expertise_area,
                'years_of_experience': t.years_of_experience,
                'created_at': t.created_at.isoformat() if t.created_at else '',
            })
        return Response(data)

    @action(detail=False, methods=['get'])
    def centers(self, request):
        from apps.centers.models import Center
        qs = Center.objects.filter(status=Center.Status.ACTIVE).values('id', 'code', 'name_bn', 'name_en')
        return Response(list(qs))

    @action(detail=False, methods=['get'])
    def courses(self, request):
        from apps.courses.models import Course
        center_id = request.query_params.get('center_id')
        qs = Course.objects.filter(status=Course.Status.ACTIVE)
        if center_id:
            qs = qs.filter(batch__center_id=center_id).distinct()
        return Response(list(qs.values('id', 'code', 'name_bn', 'name_en')))


class HOTrainerMappingViewSet(viewsets.ModelViewSet):
    queryset = TrainerMapping.objects.select_related('trainer', 'center', 'course', 'approved_by').all()
    serializer_class = TrainerMappingSerializer
    permission_classes = [permissions.IsAuthenticated, IsHeadOffice]
    filter_backends = (DjangoFilterBackend, drf_filters.SearchFilter)
    filterset_fields = ('trainer', 'center', 'course', 'status', 'is_primary')

    @action(detail=True, methods=['post'])
    def approve(self, request, pk=None):
        mapping = self.get_object()
        mapping.status = 'active'
        mapping.approved_by = request.user
        mapping.approved_at = timezone.now()
        mapping.save(update_fields=['status', 'approved_by', 'approved_at'])
        ActionLog.objects.create(
            user=request.user,
            action='approved trainer mapping',
            target_type='TrainerMapping',
            target_id=str(mapping.id),
            description=f'Approved mapping {mapping.trainer.trainer_no} → {mapping.center.code}',
            ip_address=request.META.get('REMOTE_ADDR', ''),
        )
        return Response(TrainerMappingSerializer(mapping).data)
